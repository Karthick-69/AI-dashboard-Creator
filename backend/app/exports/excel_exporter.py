import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.core.session import DatasetSession
from app.data_processing.profiler import profile_dataset

def generate_excel_workbook(session: DatasetSession) -> bytes:
    """Creates a multi-tab formatted Excel workbook containing dataset, profiling, KPIs, and audit log."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Sheet 1: Cleaned Dataset
    ws1 = wb.create_sheet(title="Cleaned Dataset")
    headers = list(session.df.columns)
    ws1.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in session.df.itertuples(index=False):
        clean_row = [None if str(v) == "nan" else v for v in row]
        ws1.append(clean_row)

    # Sheet 2: Data Profiling
    ws2 = wb.create_sheet(title="Data Quality & Profiling")
    prof_headers = ["Column Name", "Data Type", "General Type", "Null Count", "Null %", "Unique Values", "Outliers", "Min", "Mean", "Max"]
    ws2.append(prof_headers)
    for col_num in range(1, len(prof_headers) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    profile = profile_dataset(session.df, session.session_id)
    for col in profile.columns:
        num_st = col.numeric_stats
        ws2.append([
            col.name,
            col.dtype,
            col.general_type,
            col.null_count,
            f"{col.null_percentage}%",
            col.unique_count,
            col.outlier_count,
            num_st.min if num_st else "-",
            num_st.mean if num_st else "-",
            num_st.max if num_st else "-"
        ])

    # Sheet 3: KPI Metrics
    ws3 = wb.create_sheet(title="KPI Metrics")
    kpi_headers = ["Metric Title", "Value", "Aggregation", "Format"]
    ws3.append(kpi_headers)
    for col_num in range(1, len(kpi_headers) + 1):
        cell = ws3.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
    
    if session.dashboard_config and "widgets" in session.dashboard_config:
        for w in session.dashboard_config["widgets"]:
            if w.get("type") == "kpi":
                ws3.append([
                    w.get("title", ""),
                    w.get("data", ""),
                    w.get("aggregation", "sum").upper(),
                    w.get("format", "number")
                ])
    else:
        ws3.append(["Total Records", len(session.df), "COUNT", "number"])
        ws3.append(["Total Columns", len(session.df.columns), "COUNT", "number"])

    # Sheet 4: Cleaning Audit Trail
    ws4 = wb.create_sheet(title="Cleaning Audit Trail")
    audit_headers = ["Operation", "Description", "Timestamp", "Rows Before", "Rows After"]
    ws4.append(audit_headers)
    for col_num in range(1, len(audit_headers) + 1):
        cell = ws4.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    for log in session.audit_log:
        ws4.append([
            log.get("operation", ""),
            log.get("description", ""),
            log.get("timestamp", ""),
            log.get("rows_before", 0),
            log.get("rows_after", 0)
        ])

    # Auto-adjust column widths
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col[:100]:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
